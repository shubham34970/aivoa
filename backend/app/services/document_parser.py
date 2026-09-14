import email
from email import policy
import io
import os
from typing import Tuple

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

def extract_text_from_file(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Extracts plain text and format metadata from uploaded file bytes.
    Supported: PDF, DOCX, EML, TXT, MD, CSV, JSON
    """
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".pdf":
        return parse_pdf(file_bytes), "pdf"
    elif ext in [".docx", ".doc"]:
        return parse_docx(file_bytes), "docx"
    elif ext == ".eml":
        return parse_eml(file_bytes), "eml"
    else:
        # Default plain text decoding
        try:
            return file_bytes.decode("utf-8"), "txt"
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("latin-1"), "txt"
            except Exception:
                return str(file_bytes), "txt"

def parse_pdf(file_bytes: bytes) -> str:
    if PdfReader is None:
        return "PDF reader not available"
    
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts)
    except Exception as e:
        return f"Error reading PDF: {str(e)}"

def parse_docx(file_bytes: bytes) -> str:
    if docx is None:
        return "Docx reader not available"
    
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                text_parts.append(" | ".join([cell.text.strip() for cell in row.cells]))
        return "\n".join(text_parts)
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"

def parse_eml(file_bytes: bytes) -> str:
    try:
        msg = email.message_from_bytes(file_bytes, policy=policy.default)
        headers = []
        if msg["from"]: headers.append(f"From: {msg['from']}")
        if msg["to"]: headers.append(f"To: {msg['to']}")
        if msg["date"]: headers.append(f"Date: {msg['date']}")
        if msg["subject"]: headers.append(f"Subject: {msg['subject']}")
        
        header_str = "\n".join(headers)
        
        body_parts = []
        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                if content_type == "text/plain":
                    payload = part.get_payload(decode=True)
                    if payload:
                        body_parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="ignore"))
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                body_parts.append(payload.decode(msg.get_content_charset() or "utf-8", errors="ignore"))
                
        body_str = "\n".join(body_parts) if body_parts else str(msg.get_payload())
        return f"{header_str}\n\n--- EMAIL BODY ---\n{body_str}"
    except Exception as e:
        return f"Error parsing EML: {str(e)}"
